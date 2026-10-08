"""Proyectos independientes, listas, evidencias y exportación."""
import hashlib
import json
import os
import re
import shutil
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from generar_informe import generate, tabular

SCOPES = ('Red subterránea', 'Red aérea', 'Subestación tipo poste', 'Usos finales residenciales')
RESULTS = ('Pendiente', 'Cumple', 'No cumple', 'No aplica', 'Evidencia insuficiente')


def new_project():
    return dict(schema_version=1, project='', inspection='', address='', regulation='RETIE 2013',
                transition='', drive_folder='', template='', scopes=[], requirements=[])


def import_checklist(project, scope, filename):
    if scope not in SCOPES:
        raise ValueError('Selecciona un alcance válido.')
    book = load_workbook(filename, data_only=True)
    additions = []
    for sheet in book:
        # Identificar columnas por encabezados; no importar marcas de otro inspector.
        columns = None
        for cells in sheet:
            headers = {str(c.value or '').strip().upper(): c.column for c in cells}
            if 'ACTIVIDADES' in headers:
                reference = next((col for label, col in headers.items() if label.startswith('REFERENCIA')), None)
                if reference:
                    columns = (cells[0].row, headers['ACTIVIDADES'], reference)
                    break
        for row in sheet:
            if columns:
                header_row, activity_col, reference_col = columns
                if row[0].row <= header_row:
                    continue
                number, activity, reference = row[0].value, row[activity_col - 1].value, row[reference_col - 1].value
                if not isinstance(activity, str) or not activity.strip():
                    continue
                if number is not None and not isinstance(number, (int, float)):
                    continue
                item = str(number) if number is not None else f'Sin número (fila {row[0].row})'
                key = f'{scope}|{sheet.title}|fila-{row[0].row}'
            else:
                values = [c.value for c in row if c.value is not None]
                if len(values) < 3 or not isinstance(values[0], (int, float)) or not isinstance(values[1], str):
                    continue
                number, activity, reference = values[:3]
                item = str(number)
                key = f'{scope}|{sheet.title}|{number}'
            if any(r['key'] == key for r in project['requirements']) or any(r['key'] == key for r in additions):
                raise ValueError('Esta lista se superpone con requisitos importados. Crea un proyecto nuevo o usa otro alcance.')
            additions.append(dict(key=key, scope=scope, item=item, activity=activity,
                                  reference=str(reference or 'Pendiente de validar'), list_name=Path(filename).name,
                                  result='Pendiente', observation='', documents='', photos=[],
                                  reference_validated=False))
    if not additions:
        raise ValueError('No se encontraron requisitos en esa lista. Selecciona un formato de verificación compatible.')
    project['requirements'].extend(additions)
    return len(additions)


def save_project(project, filename):
    target = Path(filename)
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        try:
            previous = json.loads(target.read_text(encoding='utf-8'))
        except (ValueError, OSError):
            raise ValueError('El destino ya existe y no es un proyecto RETIE. Elige otro nombre.') from None
        if previous.get('schema_version') != 1:
            raise ValueError('No se reemplazan archivos ajenos a los proyectos de esta aplicación.')
        shutil.copy2(target, target.with_suffix(target.suffix + '.bak'))
    # Cada proyecto tiene su propia carpeta de evidencias.
    assets = target.parent / (target.stem + '_evidencias')
    saved = json.loads(json.dumps(project))
    for requirement in saved['requirements']:
        photos = []
        for raw in requirement['photos']:
            source = Path(raw)
            if not source.is_file():
                raise ValueError(f'No existe la foto: {source.name}')
            with source.open('rb') as file:
                digest = hashlib.file_digest(file, 'sha256').hexdigest()[:16]
            assets.mkdir(exist_ok=True)
            dest = assets / (digest + source.suffix.lower())
            if not dest.exists():
                shutil.copy2(source, dest)
            photos.append(str(dest.relative_to(target.parent)))
        requirement['photos'] = photos
    temporary = target.with_name(target.name + '.' + uuid4().hex + '.tmp')
    try:
        temporary.write_text(json.dumps(saved, ensure_ascii=False, indent=2), encoding='utf-8')
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def load_project(filename):
    target = Path(filename)
    project = json.loads(target.read_text(encoding='utf-8'))
    if project.get('schema_version') != 1 or not isinstance(project.get('requirements'), list):
        raise ValueError('No es un proyecto de esta aplicación.')
    for row in project['requirements']:
        row['photos'] = [str((target.parent / p).resolve()) if not Path(p).is_absolute() else p for p in row['photos']]
    return project


def validate_review(row):
    if row['result'] not in RESULTS:
        raise ValueError('Resultado de revisión inválido.')
    if row['result'] == 'No cumple':
        if not row['observation'].strip().lower().startswith('no se evidencia'):
            raise ValueError('Redacta la NC comenzando por «No se evidencia…».')
        if not row['reference'].strip():
            raise ValueError('Indica el numeral o escribe «Pendiente de validar».')
    return row


def report_case(project):
    for field in ('project', 'inspection', 'address', 'template'):
        if not project[field].strip():
            raise ValueError('Completa nombre, código, dirección y formato F-GI-51.')
    if not project['scopes']:
        raise ValueError('Selecciona al menos un alcance.')
    findings = []
    for row in project['requirements']:
        if row['scope'] not in project['scopes']:
            continue
        validate_review(row)
        if row['result'] == 'No cumple':
            findings.append(dict(id=f'NC-{len(findings)+1:02d}', scope=row['scope'],
                                 description=row['observation'], reference=row['reference'],
                                 reference_validated=row['reference_validated'], status='Pendiente de validación final',
                                 documents=[dict(name=row['documents'])] if row['documents'] else [],
                                 photos=[dict(name=Path(p).name, path=p) for p in row['photos']]))
    return dict(project=project['project'], inspection=project['inspection'], address=project['address'],
                regulation=project['regulation'] + '; transición: ' + (project['transition'] or 'por revisar'),
                scope=', '.join(project['scopes']), findings=findings)


def export_project(project, folder):
    case = report_case(project)
    if not case['findings']:
        raise ValueError('No hay NC registradas en los alcances seleccionados. Esto no significa que la instalación cumpla.')
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    name = re.sub(r'[^\w-]+', '_', project['inspection']).strip('_')[:60] or 'Informe'
    # Salidas con nombre único para conservar todas las revisiones anteriores.
    name += '_' + uuid4().hex[:8]
    data = folder / (name + '_caso.json')
    data.write_text(json.dumps(case, ensure_ascii=False, indent=2), encoding='utf-8')
    report = folder / (name + '_NC.xlsx')
    generate(data, project['template'], report)
    matrix = Workbook()
    matrix.remove(matrix.active)
    rows = [[r['scope'], r['list_name'], r['item'], r['activity'], r['reference'], r['result'],
             r['observation'], r['documents'], '\n'.join(Path(p).name for p in r['photos']),
             'Sí' if r['reference_validated'] else 'Pendiente']
            for r in project['requirements'] if r['scope'] in project['scopes']]
    tabular(matrix, 'Revisión por requisito', ['Alcance', 'Lista', 'Ítem', 'Actividad', 'Numeral',
             'Resultado', 'Descripción NC o comentario', 'Documento y página', 'Fotos', 'Numeral validado'], rows)
    for sh in matrix:
        for row in sh:
            for c in row:
                if c.data_type == 'f':
                    c.data_type = 's'
    destination = folder / (name + '_MATRIZ.xlsx')
    matrix.save(destination)
    return [report, destination]
