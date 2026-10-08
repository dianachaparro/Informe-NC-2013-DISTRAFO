import hashlib
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from PIL import Image

from generar_informe import generate


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.template = self.root / 'plantilla.xlsx'
        w = Workbook()
        s = w.active
        s.title = 'Constructor'
        s['B7'] = 'Descripción de la no conformidad detectada '
        s['B9'] = 'HALLAZGO DE OTRA OBRA'
        s['D2'] = 'OBRA ANTIGUA'
        s['K137'] = '=1+1'
        w.save(self.template)
        Image.new('RGB', (60, 40), 'white').save(self.root / 'foto.jpg')
        self.case = {'project': 'Prueba', 'inspection': 'TEST-1', 'address': 'Dirección',
                     'regulation': 'Por validar', 'scope': 'Prueba', 'findings': []}
        self.json = self.root / 'caso.json'
        self.output = self.root / 'salida.xlsx'

    def save_case(self):
        self.json.write_text(json.dumps(self.case))

    def finding(self, index):
        return {'id': f'TEST-{index}', 'description': 'Descripción nueva',
                'reference': 'Referencia por validar', 'status': 'Pendiente de validación',
                'photos': [{'name': 'foto.jpg', 'path': 'foto.jpg', 'url': 'https://example.com/foto'}]}

    def test_dynamic_rows_photos_and_original_unchanged(self):
        self.case['findings'] = [self.finding(i) for i in range(20)]
        self.save_case()
        original = hashlib.sha256(self.template.read_bytes()).hexdigest()
        self.assertEqual(generate(self.json, self.template, self.output), 20)
        w = load_workbook(self.output)
        s = w['Constructor']
        self.assertEqual(s['B28'].value, 'Descripción nueva')
        self.assertIn('pendiente', s['K30'].value)
        self.assertEqual(len(s._images), 20)
        self.assertEqual(len(w['Evidencia 20']._images), 1)
        values = [str(c.value) for sh in w for row in sh for c in row if c.value is not None]
        self.assertFalse(any('HALLAZGO DE OTRA OBRA' in x or 'OBRA ANTIGUA' in x for x in values))
        self.assertFalse(any(c.data_type == 'f' for sh in w for row in sh for c in row))
        self.assertIsNone(s['R9'].value)
        self.assertIsNone(s['S9'].value)
        self.assertEqual(original, hashlib.sha256(self.template.read_bytes()).hexdigest())

    def test_missing_evidence_does_not_silently_generate(self):
        f = self.finding(1)
        f['photos'][0]['path'] = 'ausente.jpg'
        self.case['findings'] = [f]
        self.save_case()
        with self.assertRaisesRegex(ValueError, 'No existe la foto'):
            generate(self.json, self.template, self.output)
        self.assertFalse(self.output.exists())

    def test_cannot_overwrite_original(self):
        self.save_case()
        with self.assertRaisesRegex(ValueError, 'distinta'):
            generate(self.json, self.template, self.template)


if __name__ == '__main__':
    unittest.main()
