"""Genera un borrador F-GI-51 trazable desde hallazgos revisables en JSON.

No interpreta reglamentos, no determina conformidad y no escribe en la plantilla.
"""
import argparse
from copy import copy
import json
from io import BytesIO
from PIL import Image as PillowImage, ImageOps
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.pagebreak import Break


def workbook_image(path):
    # Algunas fotos JPG son contenedores MPO; Excel requiere JPEG/PNG.
    # Se incorpora la imagen principal en memoria; el original queda intacto.
    buffer = BytesIO()
    with PillowImage.open(path) as original:
        ImageOps.exif_transpose(original).convert("RGB").save(buffer, format="JPEG", quality=95)
    buffer.seek(0)
    return Image(buffer)


def copy_row(source, target, source_row, target_row):
    for col in range(1, 20):
        a, b = source.cell(source_row, col), target.cell(target_row, col)
        b.value = a.value
        if a.has_style:
            b.font = copy(a.font)
            b.fill = copy(a.fill)
            b.border = copy(a.border)
            b.number_format = a.number_format
            b.protection = copy(a.protection)
        b.alignment = copy(a.alignment)
    target.row_dimensions[target_row].height = source.row_dimensions[source_row].height


def tabular(book, title, headings, rows):
    sheet = book.create_sheet(title)
    sheet.append(headings)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24476B')
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 42
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if isinstance(cell.value, str) and cell.value.startswith('https://'):
                cell.hyperlink = cell.value
    return sheet


def generate(case_file, template, output):
    case_file, template, output = map(Path, (case_file, template, output))
    if output.resolve() in (template.resolve(), case_file.resolve()):
        raise ValueError('La salida debe ser distinta de los archivos de entrada.')
    case = json.loads(case_file.read_text(encoding='utf-8'))
    for key in ('project', 'inspection', 'address', 'regulation', 'scope', 'findings'):
        if key not in case:
            raise ValueError(f'Falta el campo {key}.')
    seen = set()
    for finding in case['findings']:
        for field in ('id', 'description', 'reference', 'status'):
            if not finding.get(field):
                raise ValueError(f'Falta {field} en un hallazgo.')
        if finding['id'] in seen:
            raise ValueError('Identificador de hallazgo duplicado.')
        seen.add(finding['id'])
        for photo in finding.get('photos', []):
            path = Path(photo['path'])
            if not path.is_absolute():
                path = case_file.parent / path
            if not path.is_file():
                raise ValueError(f'No existe la foto {photo["name"]}.')
            photo['_local'] = path
    source_book = load_workbook(template, data_only=False, keep_links=False)
    if 'Constructor' not in source_book:
        raise ValueError('Se requiere la hoja Constructor de la plantilla F-GI-51 suministrada.')
    source = source_book['Constructor']
    if source['B7'].value != 'Descripción de la no conformidad detectada ':
        raise ValueError('La estructura del formato cambió; revisar el mapeo de campos.')
    book = Workbook()
    sheet = book.active
    sheet.title = 'Constructor'
    for key, dim in source.column_dimensions.items():
        if len(key) == 1 and key <= 'S':
            sheet.column_dimensions[key].width = dim.width
            sheet.column_dimensions[key].hidden = dim.hidden
    for row in range(1, 9):
        copy_row(source, sheet, row, row)
    for region in source.merged_cells.ranges:
        if region.max_row <= 8:
            sheet.merge_cells(str(region))
    sheet['D2'] = case['project']
    sheet['D3'] = case['address']
    sheet['D4'] = case['inspection']
    sheet.row_dimensions[3].height = 30
    evidence, annex_photos = [], []
    for index, finding in enumerate(case['findings'], 1):
        row = 8 + index
        copy_row(source, sheet, 9, row)
        for col in range(1, 20):
            sheet.cell(row, col).value = None
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        sheet.merge_cells(start_row=row, start_column=9, end_row=row, end_column=17)
        sheet.cell(row, 1, index)
        sheet.cell(row, 2, finding['description'])
        sheet.cell(row, 7, finding.get('scope', ''))
        sheet.cell(row, 8, finding['reference'])
        photos = finding.get('photos', [])
        docs = finding.get('documents', [])
        labels = [d['name'] + (f' (p. {d["page"]})' if d.get('page') else '') for d in docs]
        labels += [p['name'] for p in photos]
        sheet.cell(row, 9, '\n'.join(labels) if labels else 'Sin fotografía aplicable / evidencia pendiente.')
        sheet.cell(row, 9).alignment = Alignment(wrap_text=True, vertical='bottom')
        sheet.row_dimensions[row].height = 235 if photos else 150
        for col in range(1, 20):
            if col != 9:
                sheet.cell(row, col).alignment = Alignment(wrap_text=True, vertical='top')
        if photos:
            image = workbook_image(photos[0]['_local'])
            scale = min(300 / image.width, 185 / image.height)
            image.width *= scale
            image.height *= scale
            sheet.add_image(image, f'I{row}')
        for doc in docs:
            evidence.append([finding['id'], 'Documento', doc['name'], doc.get('page'), doc.get('url', ''), finding['status']])
        for photo in photos:
            evidence.append([finding['id'], 'Fotografía', photo['name'], None, photo.get('url', ''), finding['status']])
            annex_photos.append((finding['id'], photo))
    footer = 9 + len(case['findings'])
    for source_row, row in [(136, footer), (137, footer + 1)]:
        copy_row(source, sheet, source_row, row)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        sheet.merge_cells(start_row=row, start_column=11, end_row=row, end_column=19)
    sheet.cell(footer + 1, 11, 'Matrícula profesional: pendiente de diligenciar')
    sheet.freeze_panes = 'B9'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = f'A1:S{footer + 1}'
    sheet.print_title_rows = '1:8'
    sheet.oddHeader.center.text = 'BORRADOR — NO EMITIDO — SIN FIRMAS'
    sheet.oddFooter.center.text = 'Página &P de &N'
    for row in range(10, footer, 2):
        sheet.row_breaks.append(Break(id=row))
    tabular(book, 'Control de revisión', ['Campo', 'Detalle'], [
        ['Estado', 'Borrador preliminar. No constituye dictamen ni informe firmado.'],
        ['Reglamento', case['regulation']], ['Alcance declarado', case['scope']],
        ['Criterio', 'La ausencia de evidencia no demuestra automáticamente incumplimiento.'],
        ['Referencias', 'Numerales procedentes de listas suministradas; validar contra reglamento oficial y alcance.'],
        ['Subsanación', 'Casillas SI/NO y firmas pendientes de diligenciar.'],
        ['Referencias validadas', 'Sí' if all(f.get('reference_validated', False) for f in case['findings']) else 'Hay numerales pendientes de validar.'],
        ['Plantilla', 'Solo se copiaron estructura y estilos. Se excluyeron hallazgos, fotos y vínculos externos de otras obras.'],
    ])
    tabular(book, 'Trazabilidad', ['Hallazgo', 'Tipo', 'Archivo', 'Página PDF', 'Enlace de origen', 'Estado'], evidence)
    tabular(book, 'Puntos por verificar', ['ID', 'Observación', 'Fotografías'], [
        [n['id'], n['description'], '\n'.join(p['name'] for p in n.get('photos', []))]
        for n in case.get('review_notes', [])
    ])
    for note in case.get('review_notes', []):
        for photo in note.get('photos', []):
            path = Path(photo['path'])
            if not path.is_absolute():
                path = case_file.parent / path
            if not path.is_file():
                raise ValueError(f'No existe la foto {photo["name"]}.')
            photo['_local'] = path
            annex_photos.append((note['id'], photo))
    for index, (finding_id, photo) in enumerate(annex_photos, 1):
        annex = book.create_sheet(f'Evidencia {index:02d}')
        annex['A1'] = f'{finding_id}: {photo["name"]}'
        annex['A2'] = photo.get('url', '')
        annex['A2'].hyperlink = photo.get('url', '')
        annex['A3'] = 'Evidencia para revisión; no demuestra por sí sola conformidad o subsanación.'
        image = workbook_image(photo['_local'])
        scale = min(750 / image.width, 950 / image.height)
        image.width *= scale
        image.height *= scale
        annex.add_image(image, 'A5')
        annex.column_dimensions['A'].width = 110
        annex.page_setup.fitToWidth = 1
        annex.page_setup.fitToHeight = 1
        annex.sheet_properties.pageSetUpPr.fitToPage = True
        annex.print_area = 'A1:L65'
    for worksheet in book:
        for row in worksheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.data_type == 'f':
                    cell.data_type = 's'
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    return len(case['findings'])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--caso', required=True, type=Path)
    parser.add_argument('--plantilla', required=True, type=Path)
    parser.add_argument('--salida', required=True, type=Path)
    args = parser.parse_args()
    count = generate(args.caso, args.plantilla, args.salida)
    print(f'Borrador generado: {args.salida} ({count} hallazgos pendientes).')


if __name__ == '__main__':
    main()
